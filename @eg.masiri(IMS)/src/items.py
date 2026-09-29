from flask import Blueprint,render_template,redirect,request,flash,jsonify,send_file
from .mod import add_products_to_inventory,is_product_code_used,mark_item_sold,mark_item_refunded,remove_inventory_item,remove_inventory_items_by_code,get_inventory_stock,get_refunded_stock,get_code_assignments,reset_inventory,get_status_history,clear_status_history,get_status_history_filtered,create_sale,delete_sale
items = Blueprint("items",__name__)


# Add items

@items.route('/add_items')
def add_items():
    assignments = get_code_assignments()
    return render_template('add_items.html',assignments=assignments)

@items.route('/api/add_items',methods=["POST","GET"])
def api_add_items():
    if request.method == "POST":
        pro_code = request.values.get("product_code")
        pro_name = request.values.get("product_name")
        pro_color = request.values.get("product_color")
        pro_size = request.values.get("product_size")
        pro_price = request.values.get("product_price")
        pro_count = request.values.get("product_counter")
        if is_product_code_used(pro_code,pro_name,pro_color,pro_size):
            flash(f"code {pro_code} is already taken","danger")
            return redirect("/add_items")
        for i in range(int(pro_count)):
            add_products_to_inventory(pro_code,pro_name,pro_color,pro_size,pro_price)
    flash("items have been added successfully","success")
    return redirect("/add_items")

@items.route('/api/add_items/reset',methods=["POST"])
def api_reset_inventory():
    reset_inventory()
    flash("inventory has been reset","success")
    return redirect("/add_items")


# Orders

@items.route('/orders')
def orders():
    stock = get_inventory_stock()
    refunded = get_refunded_stock()
    return render_template("orders.html",stock=stock,refunded=refunded)

@items.route('/api/orders/add_order',methods=["POST"])
def api_add_order():
    pro_code = request.values.get("product_code")
    if pro_code and mark_item_sold(pro_code):
        flash("item marked as sold","success")
    else:
        flash("no available item with that code","danger")
    return redirect("/orders")

@items.route('/api/orders/refund_order',methods=["POST"])
def api_refund_order():
    pro_code = request.values.get("product_code")
    if pro_code and mark_item_refunded(pro_code):
        flash("item refunded","success")
    else:
        flash("no sold item with that code","danger")
    return redirect("/orders")


@items.route('/api/orders/sale',methods=["POST"])
def api_create_sale():
    pro_code = request.values.get("product_code")
    sale_amount = request.values.get("sale_amount")
    customer_name = request.values.get("customer_name")
    
    # convert empty strings to None
    sale_amount = float(sale_amount) if sale_amount else None
    customer_name = customer_name if customer_name else None
    
    if pro_code and create_sale(pro_code, sale_amount, customer_name):
        flash("item sold and recorded in sales","success")
    else:
        flash("no available item with that code","danger")
    return redirect("/orders")


@items.route('/api/sales/delete',methods=["POST"])
def api_delete_sale():
    sale_id = request.values.get("sale_id")
    if sale_id and delete_sale(int(sale_id)):
        flash("sale removed, item restored to available","success")
    else:
        flash("sale not found","danger")
    return redirect("/dashboard")


@items.route('/api/orders/remove_item',methods=["POST"])
def api_remove_item():
    pro_code = request.values.get("product_code")
    if pro_code and remove_inventory_item(pro_code):
        flash("item removed","success")
    else:
        flash("no item with that code","danger")
    return redirect("/orders")


# Add items extras

@items.route('/api/add_items/remove_all',methods=["POST"])
def api_remove_all_items():
    pro_code = request.values.get("product_code")
    removed = remove_inventory_items_by_code(pro_code) if pro_code else 0
    if removed:
        flash(f"{removed} items removed","success")
    else:
        flash("no items with that code","danger")
    return redirect("/add_items")


# History

@items.route('/history')
def history():
    start_date = request.args.get("from", "")
    end_date = request.args.get("to", "")
    if start_date or end_date:
        history_rows = get_status_history_filtered(start_date or None, end_date or None)
    else:
        history_rows = get_status_history()
    return render_template("history.html",history_rows=history_rows,start_date=start_date,end_date=end_date)

@items.route('/api/history/clear',methods=["POST"])
def api_clear_history():
    clear_status_history()
    flash("history has been cleared","success")
    return redirect("/history")


@items.route('/api/history/download')
def api_download_history():
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from .status_names import label_status
    from .mod import get_status_history_filtered

    start_date = request.args.get("from", "")
    end_date = request.args.get("to", "")
    history_rows = get_status_history_filtered(start_date or None, end_date or None)

    wb = Workbook()
    ws = wb.active
    ws.title = "History"

    headers = ["When","Code","Product","Color","Size","Old status","New status"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True,color="2B2314")
        cell.fill = PatternFill("solid",fgColor="D9A441")

    for changed_at,code,name,color,size,old_status,new_status in history_rows:
        ws.append([
            changed_at.strftime("%Y-%m-%d %H:%M"),
            code,
            name,
            color,
            size,
            label_status(old_status) if old_status else "added",
            label_status(new_status),
        ])

    # readable column widths
    widths = {"A":17,"B":8,"C":22,"D":12,"E":8,"F":13,"G":13}
    for column,width in widths.items():
        ws.column_dimensions[column].width = width

    out = BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name="history.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )